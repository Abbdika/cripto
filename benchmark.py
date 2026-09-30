"""
benchmark.py - menjalankan semua pengujian wajib Topik A dan
menyimpan hasilnya ke results/ (XLSX + grafik PNG).

Jalankan:  python benchmark.py
"""
import io
import math
import os
import secrets
import time
from collections import Counter

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw

import crypto_core as cc

OUT = "results"
PW = secrets.token_urlsafe(16)  
ALGS = list(cc.ALG_ID)
os.makedirs(OUT, exist_ok=True)


def bikin_png():
    img = Image.new("RGB", (200, 150), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 120, 100], fill="navy")
    d.ellipse([90, 60, 180, 140], fill="crimson")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def bikin_pdf():
    return (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\n"
            b"trailer<</Root 1 0 R>>\n%%EOF\n")


def shannon(data: bytes) -> float:
    n, c = len(data), Counter(data)
    return -sum(v / n * math.log2(v / n) for v in c.values())


def bitdiff(a: bytes, b: bytes) -> int:
    return (int.from_bytes(a, "big") ^ int.from_bytes(b, "big")).bit_count()


def flip_bit(data: bytes, pos: int) -> bytes:
    b = bytearray(data)
    b[pos // 8] ^= 1 << (pos % 8)
    return bytes(b)


def berkas_asli():
    """Ambil semua PDF/gambar asli dari folder data_uji/."""
    hasil, d = [], "data_uji"
    if os.path.isdir(d):
        for n in sorted(os.listdir(d)):
            if n.lower().endswith((".pdf", ".png", ".jpg", ".jpeg")):
                with open(os.path.join(d, n), "rb") as f:
                    hasil.append((f"Berkas asli: {n}", f.read()))
    if not hasil:
        print("  (data_uji/ kosong: pakai PNG/PDF buatan skrip. Taruh berkas asli di sana!)")
    return hasil


def uji_kebenaran():
    inputs = [
        ("Teks pendek", "Halo dunia".encode()),
        ("Teks Unicode", "Keamanan Informasi 🔐 Unsil Tasikmalaya".encode()),
        ("1 byte", b"A"),
        ("Data kosong", b""),
        ("Biner acak 1 KB", os.urandom(1024)),
        ("Biner acak 100 KB", os.urandom(100 * 1024)),
        ("Nol semua 1 MB", bytes(1024 * 1024)),
        ("Teks berulang 200 KB", ("lorem ipsum dolor sit amet " * 8000).encode()),
        ("Berkas gambar PNG", bikin_png()),
        ("Berkas PDF", bikin_pdf()),
    ] + berkas_asli()
    rows = []
    for nama, data in inputs:
        for alg in ALGS:
            blob = cc.encrypt(data, PW, alg)
            ok = cc.decrypt(blob, PW) == data
            try:
                cc.decrypt(blob, PW + "x"); salah = "DITERIMA (gagal!)"
            except cc.DecryptionError:
                salah = "Ditolak"
            rusak = bytearray(blob); rusak[len(rusak) // 2 if len(rusak) > 40 else -1] ^= 1
            try:
                cc.decrypt(bytes(rusak), PW); ubah = "DITERIMA (gagal!)"
            except cc.DecryptionError:
                ubah = "Ditolak"
            rows.append([nama, alg, len(data), len(blob), "Berhasil" if ok else "GAGAL", salah, ubah])
    return rows


def uji_waktu(reps=5):
    sizes = {"1 KB": 1024, "1 MB": 1024 ** 2, "10 MB": 10 * 1024 ** 2}
    rows = []
    for label, n in sizes.items():
        data = os.urandom(n)
        for alg in ALGS:
            te, td = [], []
            for _ in range(reps):
                t = time.perf_counter(); blob = cc.encrypt(data, PW, alg); te.append(time.perf_counter() - t)
                t = time.perf_counter(); cc.decrypt(blob, PW); td.append(time.perf_counter() - t)
            key, nonce = os.urandom(32), os.urandom(12)
            t = time.perf_counter()
            for _ in range(reps):
                cc.encrypt_raw(alg, key, nonce, data)
            raw = (time.perf_counter() - t) / reps
            rows.append([label, alg, round(sum(te) / reps * 1000, 2), round(sum(td) / reps * 1000, 2),
                         round(raw * 1000, 3), round(n / 1024 ** 2 / raw, 1)])
    kdf = []
    for _ in range(reps):
        t = time.perf_counter(); cc.derive_key(PW, os.urandom(16)); kdf.append(time.perf_counter() - t)
    return rows, round(sum(kdf) / reps * 1000, 2)


def uji_avalanche(trials=300):
    rows = []
    pt = os.urandom(1024)
    for alg in ALGS:
        key, nonce = os.urandom(32), os.urandom(12)
        base = cc.encrypt_raw(alg, key, nonce, pt)
        pl_body, pl_tag, k_body, k_tag = [], [], [], []
        for _ in range(trials):
            c2 = cc.encrypt_raw(alg, key, nonce, flip_bit(pt, int.from_bytes(os.urandom(2), "big") % (len(pt) * 8)))
            pl_body.append(bitdiff(base[:-16], c2[:-16]) / ((len(base) - 16) * 8) * 100)
            pl_tag.append(bitdiff(base[-16:], c2[-16:]) / 128 * 100)
            c3 = cc.encrypt_raw(alg, flip_bit(key, int.from_bytes(os.urandom(1), "big")), nonce, pt)
            k_body.append(bitdiff(base[:-16], c3[:-16]) / ((len(base) - 16) * 8) * 100)
            k_tag.append(bitdiff(base[-16:], c3[-16:]) / 128 * 100)
        m = lambda x: round(sum(x) / len(x), 3)
        rows.append([alg, "1 bit plainteks", m(pl_body), m(pl_tag)])
        rows.append([alg, "1 bit kunci", m(k_body), m(k_tag)])
    return rows


def uji_entropi():
    text = ("Keamanan informasi adalah upaya melindungi kerahasiaan, integritas, dan ketersediaan data. " * 2500).encode()
    img = Image.new("RGB", (256, 256))
    px = img.load()
    for x in range(256):
        for y in range(256):
            px[x, y] = (x, y, (x + y) // 2)
    raw_img = img.tobytes()
    samples = {"Teks": text, "Citra (piksel mentah)": raw_img}
    rows = []
    fig, axs = plt.subplots(2, 2, figsize=(10, 6))
    for i, (nama, pt) in enumerate(samples.items()):
        ct = cc.encrypt(pt, PW, "AES-256-GCM")
        rows.append([nama, len(pt), round(shannon(pt), 4), round(shannon(ct), 4)])
        for j, (lbl, d) in enumerate((("plainteks", pt), ("cipherteks AES-GCM", ct))):
            axs[i][j].hist(list(d), bins=256, range=(0, 255), color="#2b59c3" if j == 0 else "#c33b2b")
            axs[i][j].set_title(f"{nama} - {lbl} (H={shannon(d):.3f})", fontsize=9)
    for a in axs.flat:
        a.set_xlabel("nilai byte"); a.set_ylabel("frekuensi")
    plt.tight_layout(); plt.savefig(f"{OUT}/histogram_entropi.png", dpi=130); plt.close()
    return rows


def visual_ecb():
    W = H = 256
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([30, 30, 130, 130], fill="navy")
    d.ellipse([100, 100, 220, 220], fill="crimson")
    d.polygon([(180, 20), (240, 90), (160, 90)], fill="green")
    raw = img.tobytes()
    key = os.urandom(32)
    pad = (-len(raw)) % 16
    ecb = Cipher(algorithms.AES(key), modes.ECB()).encryptor()  
    ct_ecb = (ecb.update(raw + bytes(pad)) + ecb.finalize())[:len(raw)]
    ct_gcm = cc.encrypt_raw("AES-256-GCM", key, os.urandom(12), raw)[:len(raw)]
    fig, axs = plt.subplots(1, 3, figsize=(10, 3.6))
    for a, (t, b) in zip(axs, (("Citra asli", raw), ("AES-ECB (tidak aman)", ct_ecb), ("AES-256-GCM", ct_gcm))):
        a.imshow(Image.frombytes("RGB", (W, H), b)); a.set_title(t); a.axis("off")
    plt.tight_layout(); plt.savefig(f"{OUT}/perbandingan_ecb_gcm.png", dpi=130); plt.close()


def tulis(wb, judul, header, rows):
    ws = wb.create_sheet(judul)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or "")) for c in col) + 3
    return ws


def main():
    print("[1/5] kebenaran dekripsi..."); k = uji_kebenaran()
    print("[2/5] waktu enkripsi/dekripsi..."); w, kdf = uji_waktu()
    print("[3/5] avalanche effect..."); a = uji_avalanche()
    print("[4/5] entropi & histogram..."); e = uji_entropi()
    print("[5/5] visualisasi ECB vs GCM..."); visual_ecb()

    wb = Workbook(); wb.remove(wb.active)
    tulis(wb, "Kebenaran", ["Masukan", "Algoritma", "Ukuran plain (B)", "Ukuran cipher (B)", "Dekripsi",
                            "Password salah", "1 byte diubah"], k)
    ws = tulis(wb, "Waktu", ["Ukuran", "Algoritma", "Enkripsi total (ms)", "Dekripsi total (ms)",
                            "Cipher saja (ms)", "Throughput (MB/s)"], w)
    ws.append([]); ws.append([f"Rata-rata waktu scrypt (KDF): {kdf} ms (N=2^15, r=8, p=1)"])
    tulis(wb, "Avalanche", ["Algoritma", "Perubahan", "Bit cipherteks berubah (%)", "Bit tag berubah (%)"], a)
    tulis(wb, "Entropi", ["Data", "Ukuran (B)", "Entropi plainteks (bit/byte)", "Entropi cipherteks (bit/byte)"], e)
    wb.save(f"{OUT}/hasil_pengujian.xlsx")

    print("\n== RINGKASAN ==")
    print("Kebenaran:", sum(r[4] == "Berhasil" for r in k), "/", len(k), "berhasil;",
    "semua negatif ditolak:", all(r[5] == r[6] == "Ditolak" for r in k))
    print("KDF scrypt:", kdf, "ms")
    for r in w: print("Waktu", r)
    for r in a: print("Avalanche", r)
    for r in e: print("Entropi", r)


if __name__ == "__main__":
    main()
